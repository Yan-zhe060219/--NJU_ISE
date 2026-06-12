"""预处理 inbox/ 中的 PPTX —— 使用 python-pptx 提取文本（含演讲者备注）。

用法：
  python tools/extract_pptx.py                     # 处理 inbox 下全部 PPTX
  python tools/extract_pptx.py inbox/OS            # 处理指定课程文件夹
  python tools/extract_pptx.py --force inbox/OS    # 强制重新提取（覆盖已有 .txt）
  python tools/extract_pptx.py --single inbox/OS   # 只提取下一个未处理的（供 /learn 自动调用）

输出：在 PPTX 所在目录的 .cache/ 子目录生成 .txt，与 PPTX 隔离；进度记录在 inbox/.extracted_pptx
"""

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.shapes.base import BaseShape


CACHE_DIR = ".cache"


def has_text_frame(shape: BaseShape) -> bool:
    try:
        return shape.has_text_frame
    except Exception:
        return False


def extract_text(pptx_path: Path) -> str:
    prs = Presentation(str(pptx_path))
    slides_out = []

    for i, slide in enumerate(prs.slides, 1):
        parts = [f"=== 第 {i} 页幻灯片 ==="]

        # 幻灯片正文
        for shape in slide.shapes:
            if has_text_frame(shape):
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)

        # 演讲者备注
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[演讲者备注] {notes}")

        slides_out.append("\n".join(parts))

    return "\n\n".join(slides_out)


def load_extracted(extracted_path: Path) -> set[str]:
    if extracted_path.exists():
        return {line.strip() for line in extracted_path.read_text(encoding="utf-8").splitlines() if line.strip()}
    return set()


def save_extracted(extracted_path: Path, entry: str) -> None:
    with open(extracted_path, "a", encoding="utf-8") as f:
        f.write(entry + "\n")


def process_single(pptx_path: Path, force: bool) -> str:
    cache_dir = pptx_path.parent / CACHE_DIR
    cache_dir.mkdir(parents=True, exist_ok=True)
    txt_path = cache_dir / (pptx_path.stem + ".txt")

    if txt_path.exists() and not force:
        return "跳过（已有 .txt）"

    try:
        text = extract_text(pptx_path)
        if not text.strip():
            return "警告：未提取到文本（可能全是图片/图表）"
        txt_path.write_text(text, encoding="utf-8")
        return "完成"
    except Exception as e:
        return f"失败: {e}"


def collect_pptx(root: Path, pattern: str | None) -> list[Path]:
    if pattern:
        target = Path(pattern)
        if target.is_file():
            return [target]
        return sorted(target.rglob("*.pptx"))
    return sorted(root.rglob("*.pptx"))


def main():
    parser = argparse.ArgumentParser(description="从 inbox 中提取 PPTX 文本")
    parser.add_argument("target", nargs="?", default=None, help="课程文件夹或单个 PPTX")
    parser.add_argument("--force", action="store_true", help="强制重新提取，覆盖已有 .txt")
    parser.add_argument("--single", action="store_true", help="只处理下一篇未提取的文件（供 /learn 自动调用）")
    args = parser.parse_args()

    project_root = Path(__file__).resolve().parent.parent
    inbox_root = project_root / "inbox"
    extracted_file = inbox_root / ".extracted_pptx"

    if args.target:
        candidate = Path(args.target)
        if not candidate.is_absolute():
            candidate = project_root / candidate
    else:
        candidate = inbox_root

    pptx_list = collect_pptx(candidate if candidate.is_dir() else inbox_root,
                             str(candidate) if candidate.is_file() else None)

    already = load_extracted(extracted_file)

    if args.single:
        for pptx_file in pptx_list:
            if pptx_file.name not in already:
                pptx_list = [pptx_file]
                break
        else:
            print("所有 PPTX 文件均已提取，无需处理。")
            return

    to_process = [p for p in pptx_list if p.name not in already] if not args.force else pptx_list

    if not to_process:
        print("没有新的 PPTX 需要提取。")
        return

    print(f"共 {len(pptx_list)} 个 PPTX，待提取 {len(to_process)} 个\n")

    for i, pptx_file in enumerate(to_process, 1):
        rel = pptx_file.relative_to(project_root)
        print(f"[{i}/{len(to_process)}] {rel} ... ", end="", flush=True)
        result = process_single(pptx_file, args.force)
        print(result)
        save_extracted(extracted_file, pptx_file.name)

    print(f"\n全部完成。提取的文本已保存为同名 .txt 文件。")


if __name__ == "__main__":
    main()
